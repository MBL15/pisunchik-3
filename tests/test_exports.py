"""Выгрузка документов: рендереры xlsx/docx/pdf и построители документов поверх данных сервисов."""
import io
import unittest
from datetime import date
from unittest import mock

from docx import Document
from openpyxl import load_workbook

from gateway_service import exports, reports
from gateway_service.exports import Doc, ExportError, Heading, Para, Signatures, Table
from sportsman_service.headcount import build_report
from tests.test_headcount import A, B, C, TODAY

SAMPLE = Doc(title="Тестовый документ", filename="Тест", header_lines=["Шапка организации"], subtitle="подзаголовок",
             meta=[("Ключ", "значение")],
             blocks=[Heading("Раздел"), Para("Абзац с символами & < > «кавычки»"),
                     Table(["Имя", "Число", "Дробное"], [["Иванов", 3, 2.5], ["Петров", 10, None]], title="Таблица",
                           widths=[2, 1, 1], align=["left", "right", "right"], highlight={1}),
                     Signatures([("Руководитель", "Сидоров С.С.")])], footer="подвал")


def docx_text(content: bytes) -> str:
    d = Document(io.BytesIO(content))
    return "\n".join([p.text for p in d.paragraphs] + [c.text for t in d.tables for r in t.rows for c in r.cells])


def xlsx_cells(content: bytes) -> list:
    ws = load_workbook(io.BytesIO(content)).active
    return [c for row in ws.iter_rows(values_only=True) for c in row if c is not None]


class RenderersTest(unittest.TestCase):
    def test_xlsx_keeps_numbers_numeric_and_cyrillic_text(self):
        cells = xlsx_cells(exports.render(SAMPLE, "xlsx").content)
        for expected in ("Тестовый документ", "Шапка организации", "Иванов", 3, 2.5, "Сидоров С.С.", "подвал"):
            self.assertIn(expected, cells)

    def test_xlsx_never_turns_registry_text_into_a_formula(self):
        evil = '=HYPERLINK("http://example.com","клик")'
        doc = Doc(title="Тест", filename="t", meta=[("Ключ", "=1+1")], blocks=[Table(["Имя"], [[evil]])])
        ws = load_workbook(io.BytesIO(exports.render(doc, "xlsx").content)).active
        cells = {c.value: c.data_type for row in ws.iter_rows() for c in row if c.value is not None}
        self.assertEqual((cells[evil], cells["=1+1"]), ("s", "s"))

    def test_docx_contains_headings_table_and_signature(self):
        text = docx_text(exports.render(SAMPLE, "docx").content)
        for expected in ("Тестовый документ", "Раздел", "«кавычки»", "Иванов", "2,5", "Руководитель", "Сидоров С.С."):
            self.assertIn(expected, text)

    def test_pdf_is_a_complete_document_with_embedded_cyrillic_font(self):
        pdf = exports.render(SAMPLE, "pdf")
        self.assertEqual((pdf.media_type, pdf.filename), ("application/pdf", "Тест.pdf"))
        self.assertTrue(pdf.content.startswith(b"%PDF") and pdf.content.rstrip().endswith(b"%%EOF"))
        self.assertIn(b"/FontFile2", pdf.content)

    def test_unknown_format_and_missing_font_are_reported(self):
        with self.assertRaises(ExportError):
            exports.render(SAMPLE, "odt")
        with mock.patch.object(exports, "_fonts", None), mock.patch.object(exports, "_FONT_CANDIDATES", []):
            with self.assertRaisesRegex(ExportError, "шрифт"):
                exports.render(SAMPLE, "pdf")

    def test_download_header_has_ascii_fallback_and_utf8_name(self):
        h = exports.content_disposition("Протокол ОРБ-ALG-2026-01.pdf")
        self.assertTrue(h.isascii())
        self.assertIn("filename*=UTF-8''%D0%9F%D1%80%D0%BE%D1%82%D0%BE%D0%BA%D0%BE%D0%BB", h)
        self.assertIn('filename="', h)


COMPETITION = {"code": "ОРБ-ALG-2026-01", "title": "Первенство области", "plan_label": "ЕКП", "plan_number": "№ 1", "discipline_title": "Алгоритмическое",
               "level_status": "региональное", "type": "Первенство", "start_date": "2026-05-01", "end_date": "2026-05-02",
               "venue": "г. Оренбург", "format_label": "Очно", "format_note": "", "age_group": "без ограничений 12+"}


class BuildersTest(unittest.TestCase):
    def test_protocol_rows_dsq_and_signatures(self):
        entries = [{"place": 1, "status": "OK", "full_name": "Иванов Иван", "fsp_id": "FSP-56-25-00001-6", "region": "Оренбургская обл.",
                    "solved": 5, "penalty": 410, "score": 0, "dsq_reason": ""},
                   {"place": None, "status": "DSQ", "full_name": "Жуков Артём", "fsp_id": "FSP-56-25-00012-9", "region": "г. Новотроицк",
                    "solved": 0, "penalty": 0, "score": 0, "dsq_reason": "Нарушение"}]
        p = {"competition": COMPETITION, "mode": "icpc", "entries": entries, "published": False, "source": "manual",
             "updated_at": "2026-05-02T10:00:00"}
        doc = reports.protocol_doc(p, {"Главный судья": "Малахова О.С."})
        text = docx_text(exports.render(doc, "docx").content)
        for expected in ("ПРОТОКОЛ", "ЧЕРНОВИК", "DSQ", "Нарушение", "Малахова О.С.", "дисквалифицировано 1", "01.05.2026 — 02.05.2026"):
            self.assertIn(expected, text)
        self.assertEqual(exports.render(doc, "pdf").content[:4], b"%PDF")

    def test_headcount_report_has_a_table_per_cut_and_the_totals(self):
        r = build_report([A, B, C], TODAY, years=3)
        doc = reports.headcount_doc(r, "без фильтров")
        cells = xlsx_cells(exports.render(doc, "xlsx").content)
        self.assertIn("Занимающихся — всего", cells)
        self.assertEqual([c for c in cells if isinstance(c, int)][:3], [1, 2, 3])
        self.assertIn("+2", cells)
        self.assertIn("+200,0 %", cells)
        titles = [b.title for b in doc.blocks if isinstance(b, Table)]
        self.assertEqual(titles[0], "Сводка")
        self.assertIn("По дисциплинам", titles)
        self.assertIn("динамика за 2 года: 2024–2026", doc.subtitle)

    def test_team_documents(self):
        member = {"fsp_id": "FSP-56-25-00003-8", "full_name": "Соколов Дмитрий Алексеевич", "role": "MAIN", "position": 1, "points": 222.0,
                  "rank_short": "КМС", "municipality": "г. Оренбург", "organization": "ОГУ", "birth_date": "2004-02-17",
                  "coaches": ["Кузнецов Андрей Викторович"]}
        t = {"id": 1, "title": "Сборная по алгоритмике", "discipline": "ALG", "discipline_title": "Программирование алгоритмическое",
             "season": 2026, "purpose": "участия в чемпионате России", "age_group": None, "min_rank": "2", "status": "DRAFT",
             "status_label": "Состав формируется", "locked": False, "order_no": "", "order_date": None, "main_count": 1, "reserve_count": 1,
             "coaches": ["Кузнецов Андрей Викторович"], "updated_at": "2026-10-27T10:00:00",
             "members": [member, {**member, "fsp_id": "FSP-56-25-00008-7", "full_name": "Абрамов Кирилл Сергеевич", "role": "RESERVE", "position": 2}]}
        roster = docx_text(exports.render(reports.team_roster_doc(t), "docx").content)
        for expected in ("СПИСОК СБОРНОЙ КОМАНДЫ", "Основной состав", "Резерв", "Соколов Дмитрий Алексеевич", "17.02.2004", "222",
                         "рейтинг сезона 2026", "разряд не ниже 2"):
            self.assertIn(expected, roster)
        draft = docx_text(exports.render(reports.team_order_doc(t), "docx").content)
        for expected in ("ПРОЕКТ", "ПРИКАЗ", "ПРИКАЗЫВАЮ:", "участия в чемпионате России", "в количестве 1 спортсменов основного состава и 1",
                         "Кузнецов Андрей Викторович", "№ ______", "Приложение к приказу", "Министерство физической культуры"):
            self.assertIn(expected, draft)
        approved = docx_text(exports.render(reports.team_order_doc({**t, "locked": True, "order_no": "412-р", "order_date": "2026-11-01"}), "docx").content)
        self.assertIn("от 01.11.2026 № 412-р", approved)
        self.assertNotIn("№ ______", approved)
        typed_with_sign = docx_text(exports.render(reports.team_order_doc({**t, "locked": True, "order_no": "№ 412-р",
                                                                           "order_date": "2026-11-01"}), "docx").content)
        self.assertIn("от 01.11.2026 № 412-р", typed_with_sign)
        self.assertNotIn("№ №", typed_with_sign)

    def test_athlete_report(self):
        person = {"fsp_id": "FSP-56-25-00008-7", "full_name": "Абрамов Кирилл Сергеевич", "short_name": "Абрамов К. С.",
                  "birth_date": "2008-05-21", "municipality": "г. Оренбург", "organization": "Лицей № 1", "discipline_labels": ["Алгоритмика"],
                  "rank": {"title": "I спортивный разряд", "expires_at": "2026-11-09"}}
        base = {"competition_code": "X", "level_label": "Региональные", "protocol_size": 12, "base_points": 40, "title": "Кубок"}
        perf = {"starts": 2, "podiums": 1, "golds": 0, "best_place": 2, "results": [
            {**base, "date": "2026-08-01", "place": 2, "status": "OK", "mode": "icpc", "solved": 6, "penalty": 500, "score": 0},
            {**base, "date": "2026-05-01", "place": None, "status": "DSQ", "mode": "score", "solved": 0, "penalty": 0, "score": 0, "base_points": 0}]}
        text = docx_text(exports.render(reports.athlete_report_doc(person, perf, date(2026, 10, 27)), "docx").content)
        for expected in ("ОТЧЁТ О СОРЕВНОВАНИЯХ СПОРТСМЕНА", "I спортивный разряд (до 09.11.2026)", "решено 6, штраф 500 мин", "дисквалификация",
                         "стартов — 2, призовых мест — 1"):
            self.assertIn(expected, text)


if __name__ == "__main__":
    unittest.main()
