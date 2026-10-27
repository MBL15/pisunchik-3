"""Выгрузка документов в .xlsx / .docx / .pdf из одного описания (`Doc`): шапка, реквизиты, заголовки, абзацы,
таблицы и подписи. Построители документов — в `reports.py`; здесь только рендеринг.

PDF строится reportlab и требует TTF-шрифт с кириллицей (Arial / DejaVu / Liberation из системы или `FSP_PDF_FONT`).
"""
import io
import math
import re
from dataclasses import dataclass, field
from pathlib import Path
from urllib.parse import quote
from xml.sax.saxutils import escape

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table as PdfTable, TableStyle

from shared import config
from shared.domain import FEDERATION_NAME

MEDIA_TYPES = {
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "pdf": "application/pdf",
}
BRAND, BRAND_TEXT, HIGHLIGHT = "402FFF", "FFFFFF", "FFF2CC"


class ExportError(Exception):
    pass


@dataclass
class Heading:
    text: str


@dataclass
class Para:
    text: str
    bold: bool = False
    align: str = "left"  # left | center | right
    small: bool = False


@dataclass
class Table:
    columns: list[str]
    rows: list[list]  # ячейки: str | int | float | None
    title: str = ""
    widths: list[float] | None = None  # относительные ширины колонок
    align: list[str] | None = None  # по колонке: left | center | right
    highlight: set[int] = field(default_factory=set)  # индексы выделяемых строк


@dataclass
class Signatures:
    items: list[tuple[str, str]]  # (должность, ФИО)


@dataclass
class Doc:
    title: str
    filename: str
    subtitle: str = ""
    header_lines: list[str] = field(default_factory=list)  # шапка организации над заголовком
    meta: list[tuple[str, str]] = field(default_factory=list)  # реквизиты «ключ — значение»
    blocks: list = field(default_factory=list)
    footer: str = ""
    landscape: bool = False
    font: str = "Times New Roman"  # шрифт .docx; в PDF — найденный системный


@dataclass
class Rendered:
    content: bytes
    media_type: str
    filename: str


def _text(v, numeric_ok: bool = False):
    if v is None:
        return ""
    if isinstance(v, float):
        return v if numeric_ok else f"{v:.2f}".rstrip("0").rstrip(".").replace(".", ",")
    return v if numeric_ok and isinstance(v, int) else str(v)


def _tables(doc: Doc) -> list[Table]:
    return [b for b in doc.blocks if isinstance(b, Table)]


# ---------------------------------------------------------------- xlsx
def to_xlsx(doc: Doc) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = re.sub(r"[\[\]:*?/\\]", " ", doc.title)[:31] or "Документ"
    ncols = max([len(t.columns) for t in _tables(doc)] + [2])
    thin = Side(style="thin", color="B0B0B8")
    box = Border(left=thin, right=thin, top=thin, bottom=thin)
    last = get_column_letter(ncols)

    widths = [10.0] * ncols
    for t in _tables(doc):
        for i, name in enumerate(t.columns):
            cells = [len(str(name)) * 0.8] + [len(str(_text(r[i]))) for r in t.rows if i < len(r)]
            widths[i] = min(60.0, max(widths[i], max(cells) + 2))
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    page_width = sum(widths)

    row = 1

    def line(text, bold=False, size=11, align="left", fill=None):
        nonlocal row
        ws.merge_cells(f"A{row}:{last}{row}")
        c = ws[f"A{row}"]
        c.value, c.font = text, Font(bold=bold, size=size)
        c.alignment = Alignment(horizontal=align, vertical="top", wrap_text=True)
        if fill:
            c.fill = PatternFill("solid", fgColor=fill)
        ws.row_dimensions[row].height = max(1, math.ceil(len(str(text)) * size / 11 / max(page_width, 20))) * (size + 5)
        row += 1

    for h in doc.header_lines:
        line(h, bold=True, size=10, align="center")
    line(doc.title, bold=True, size=14, align="center")
    if doc.subtitle:
        line(doc.subtitle, size=12, align="center")
    row += 1
    for key, value in doc.meta:
        ws[f"A{row}"].value, ws[f"A{row}"].font = key, Font(bold=True)
        ws[f"A{row}"].alignment = Alignment(vertical="top", wrap_text=True)
        if ncols > 1:
            ws.merge_cells(f"B{row}:{last}{row}")
        ws[f"B{row}"].value = value
        ws[f"B{row}"].alignment = Alignment(vertical="top", wrap_text=True)
        row += 1
    for b in doc.blocks:
        row += 1
        if isinstance(b, Heading):
            line(b.text, bold=True, size=12)
        elif isinstance(b, Para):
            line(b.text, bold=b.bold, size=9 if b.small else 11, align=b.align)
        elif isinstance(b, Table):
            if b.title:
                line(b.title, bold=True, size=12)
            for i, name in enumerate(b.columns, start=1):
                c = ws.cell(row=row, column=i, value=name)
                c.font, c.border = Font(bold=True, color=BRAND_TEXT), box
                c.fill = PatternFill("solid", fgColor=BRAND)
                c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            ws.row_dimensions[row].height = 32
            row += 1
            for ri, r in enumerate(b.rows):
                for i in range(len(b.columns)):
                    c = ws.cell(row=row, column=i + 1, value=_text(r[i] if i < len(r) else None, numeric_ok=True))
                    c.border = box
                    c.alignment = Alignment(horizontal=(b.align[i] if b.align else "left"), vertical="top", wrap_text=True)
                    if ri in b.highlight:
                        c.fill = PatternFill("solid", fgColor=HIGHLIGHT)
                row += 1
        elif isinstance(b, Signatures):
            for position, name in b.items:
                ws[f"A{row}"].value = position
                ws[f"A{row}"].alignment = Alignment(wrap_text=True, vertical="bottom")
                if ncols > 2:
                    ws[f"{get_column_letter(max(ncols - 1, 2))}{row}"].value = "________________"
                    ws[f"{last}{row}"].value = name
                ws.row_dimensions[row].height = 30
                row += 1
    if doc.footer:
        row += 1
        line(doc.footer, size=9)

    for cells in ws.iter_rows():  # текст из реестра («=HYPERLINK(…)») не должен становиться формулой
        for c in cells:
            if c.data_type == "f":
                c.data_type = "s"
    ws.page_setup.orientation = "landscape" if doc.landscape else "portrait"
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------- docx
_DOCX_ALIGN = {"left": WD_ALIGN_PARAGRAPH.LEFT, "center": WD_ALIGN_PARAGRAPH.CENTER, "right": WD_ALIGN_PARAGRAPH.RIGHT}


def _shade(cell, color: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), color)
    tc_pr.append(shd)


def _run(p, text, doc: Doc, size=11, bold=False, italic=False, color=None):
    run = p.add_run(str(text))
    run.bold, run.italic = bold, italic
    run.font.size, run.font.name = Pt(size), doc.font
    run._element.rPr.rFonts.set(qn("w:eastAsia"), doc.font)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    return run


def _para(d, doc: Doc, text="", size=11, bold=False, italic=False, align="left", before=0, after=4):
    p = d.add_paragraph()
    p.alignment = _DOCX_ALIGN[align]
    p.paragraph_format.space_before, p.paragraph_format.space_after = Pt(before), Pt(after)
    if text != "":
        _run(p, text, doc, size, bold, italic)
    return p


def to_docx(doc: Doc) -> bytes:
    d = Document()
    section = d.sections[0]
    if doc.landscape:
        section.orientation = WD_ORIENT.LANDSCAPE
        section.page_width, section.page_height = Mm(297), Mm(210)
    else:
        section.page_width, section.page_height = Mm(210), Mm(297)
    section.left_margin = section.right_margin = Mm(18 if not doc.landscape else 15)
    section.top_margin = section.bottom_margin = Mm(15)
    usable = (section.page_width - section.left_margin - section.right_margin)
    normal = d.styles["Normal"]
    normal.font.name, normal.font.size = doc.font, Pt(11)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), doc.font)

    for h in doc.header_lines:
        _para(d, doc, h, 10, bold=True, align="center", after=0)
    _para(d, doc, doc.title, 15, bold=True, align="center", before=8, after=2)
    if doc.subtitle:
        _para(d, doc, doc.subtitle, 12, align="center", after=6)
    for key, value in doc.meta:
        p = _para(d, doc, after=1)
        _run(p, f"{key}: ", doc, 10.5, bold=True)
        _run(p, value, doc, 10.5)

    for b in doc.blocks:
        if isinstance(b, Heading):
            _para(d, doc, b.text, 12, bold=True, before=10, after=4)
        elif isinstance(b, Para):
            _para(d, doc, b.text, 9 if b.small else 11, bold=b.bold, align=b.align, after=6)
        elif isinstance(b, Table):
            if b.title:
                _para(d, doc, b.title, 11.5, bold=True, before=10, after=4)
            tbl = d.add_table(rows=1, cols=len(b.columns))
            tbl.style, tbl.alignment = "Table Grid", WD_TABLE_ALIGNMENT.CENTER
            weights = b.widths or [1] * len(b.columns)
            for i, name in enumerate(b.columns):
                cell = tbl.rows[0].cells[i]
                cell.width = int(usable * weights[i] / sum(weights))
                _shade(cell, BRAND)
                p = cell.paragraphs[0]
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                _run(p, name, doc, 9.5, bold=True, color=BRAND_TEXT)
            for ri, r in enumerate(b.rows):
                cells = tbl.add_row().cells
                for i in range(len(b.columns)):
                    cells[i].width = int(usable * weights[i] / sum(weights))
                    p = cells[i].paragraphs[0]
                    p.alignment = _DOCX_ALIGN[b.align[i] if b.align else "left"]
                    _run(p, _text(r[i] if i < len(r) else None), doc, 9.5)
                    if ri in b.highlight:
                        _shade(cells[i], HIGHLIGHT)
            # повтор шапки на каждой странице
            trpr = tbl.rows[0]._tr.get_or_add_trPr()
            header = OxmlElement("w:tblHeader")
            header.set(qn("w:val"), "true")
            trpr.append(header)
        elif isinstance(b, Signatures):
            tbl = d.add_table(rows=0, cols=3)
            for position, name in b.items:
                cells = tbl.add_row().cells
                cells[0].width, cells[1].width, cells[2].width = int(usable * 0.45), int(usable * 0.3), int(usable * 0.25)
                _run(cells[0].paragraphs[0], position, doc, 11)
                _run(cells[1].paragraphs[0], "________________", doc, 11)
                _run(cells[2].paragraphs[0], name, doc, 11)
                cells[0].paragraphs[0].paragraph_format.space_before = Pt(18)
    if doc.footer:
        _para(d, doc, doc.footer, 9, italic=True, before=12)
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------- pdf
_FONT_CANDIDATES = [
    ("C:/Windows/Fonts/arial.ttf", "C:/Windows/Fonts/arialbd.ttf"),
    ("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf", "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf"),
    ("/usr/share/fonts/TTF/DejaVuSans.ttf", "/usr/share/fonts/TTF/DejaVuSans-Bold.ttf"),
    ("/usr/share/fonts/truetype/liberation/LiberationSans-Regular.ttf", "/usr/share/fonts/truetype/liberation/LiberationSans-Bold.ttf"),
    ("/Library/Fonts/Arial.ttf", "/Library/Fonts/Arial Bold.ttf"),
    ("/System/Library/Fonts/Supplemental/Arial.ttf", "/System/Library/Fonts/Supplemental/Arial Bold.ttf"),
]
_fonts: tuple[str, str] | None = None


def _pdf_fonts() -> tuple[str, str]:
    global _fonts
    if _fonts:
        return _fonts
    candidates = ([(config.PDF_FONT, config.PDF_FONT_BOLD or config.PDF_FONT)] if config.PDF_FONT else []) + _FONT_CANDIDATES
    for regular, bold in candidates:
        if Path(regular).is_file() and Path(bold).is_file():
            pdfmetrics.registerFont(TTFont("FSP", regular))
            pdfmetrics.registerFont(TTFont("FSP-Bold", bold))
            _fonts = ("FSP", "FSP-Bold")
            return _fonts
    raise ExportError("Для PDF не найден шрифт с кириллицей: укажите путь к TTF в FSP_PDF_FONT (и FSP_PDF_FONT_BOLD)")


def _p(text, style) -> Paragraph:
    return Paragraph(escape(str(text)).replace("\n", "<br/>"), style)


def to_pdf(doc: Doc) -> bytes:
    regular, bold = _pdf_fonts()
    base = ParagraphStyle("base", fontName=regular, fontSize=10, leading=13)

    def style(name, **kw):
        return ParagraphStyle(name, parent=base, **kw)

    s_title = style("title", fontName=bold, fontSize=15, leading=19, alignment=1, spaceBefore=6, spaceAfter=2)
    s_center = style("center", alignment=1)
    s_head = style("head", fontName=bold, fontSize=11.5, leading=15, spaceBefore=10, spaceAfter=4)
    s_cell = style("cell", fontSize=8.5, leading=10.5)
    s_cell_head = style("cellhead", fontName=bold, fontSize=8.5, leading=10.5, textColor=colors.white, alignment=1)
    align = {"left": 0, "center": 1, "right": 2}
    cell_styles = {a: ParagraphStyle(f"cell-{a}", parent=s_cell, alignment=n) for a, n in align.items()}

    buf = io.BytesIO()
    size = landscape(A4) if doc.landscape else A4
    pdf = SimpleDocTemplate(buf, pagesize=size, leftMargin=15 * mm, rightMargin=15 * mm, topMargin=14 * mm,
                            bottomMargin=16 * mm, title=doc.title, author=FEDERATION_NAME)
    usable = size[0] - 30 * mm
    story: list = []
    for h in doc.header_lines:
        story.append(_p(h, style("hdr", fontName=bold, fontSize=9, leading=11, alignment=1)))
    story.append(_p(doc.title, s_title))
    if doc.subtitle:
        story.append(_p(doc.subtitle, style("sub", fontSize=11.5, leading=15, alignment=1, spaceAfter=6)))
    if doc.meta:
        meta = PdfTable([[_p(k, style("mk", fontName=bold, fontSize=9, leading=11.5)), _p(v, style("mv", fontSize=9, leading=11.5))]
                         for k, v in doc.meta], colWidths=[usable * 0.28, usable * 0.72])
        meta.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "TOP"), ("BOTTOMPADDING", (0, 0), (-1, -1), 2),
                                  ("TOPPADDING", (0, 0), (-1, -1), 2), ("LINEBELOW", (0, 0), (-1, -1), 0.25, colors.HexColor("#D9D9D9"))]))
        story += [Spacer(1, 6), meta]

    for b in doc.blocks:
        if isinstance(b, Heading):
            story.append(_p(b.text, s_head))
        elif isinstance(b, Para):
            story.append(_p(b.text, style("para", fontName=bold if b.bold else regular, alignment=align[b.align],
                                          fontSize=8 if b.small else 10, leading=10 if b.small else 14, spaceAfter=6)))
        elif isinstance(b, Table):
            if b.title:
                story.append(_p(b.title, s_head))
            weights = b.widths or [1] * len(b.columns)
            data = [[_p(c, s_cell_head) for c in b.columns]]
            for r in b.rows:
                data.append([_p(_text(r[i] if i < len(r) else None), cell_styles[b.align[i] if b.align else "left"])
                             for i in range(len(b.columns))])
            t = PdfTable(data, colWidths=[usable * w / sum(weights) for w in weights], repeatRows=1)
            ts = [("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#" + BRAND)), ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#9A9AA3")),
                  ("VALIGN", (0, 0), (-1, -1), "TOP"), ("TOPPADDING", (0, 0), (-1, -1), 3), ("BOTTOMPADDING", (0, 0), (-1, -1), 3)]
            ts += [("BACKGROUND", (0, ri + 1), (-1, ri + 1), colors.HexColor("#" + HIGHLIGHT)) for ri in b.highlight]
            t.setStyle(TableStyle(ts))
            story += [Spacer(1, 4), t]
        elif isinstance(b, Signatures):
            sig = PdfTable([[_p(position, base), "", _p(name, base)] for position, name in b.items],
                           colWidths=[usable * 0.45, usable * 0.3, usable * 0.25], rowHeights=[26] * len(b.items))
            sig.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "BOTTOM"), ("LINEBELOW", (1, 0), (1, -1), 0.6, colors.black)]))
            story += [Spacer(1, 14), sig]
    if doc.footer:
        story += [Spacer(1, 12), _p(doc.footer, style("foot", fontSize=8, leading=10, textColor=colors.HexColor("#5F5F66")))]

    def numbering(canvas, d):
        canvas.saveState()
        canvas.setFont(regular, 8)
        canvas.setFillColor(colors.HexColor("#5F5F66"))
        canvas.drawRightString(size[0] - 15 * mm, 9 * mm, f"стр. {d.page}")
        canvas.restoreState()

    pdf.build(story, onFirstPage=numbering, onLaterPages=numbering)
    return buf.getvalue()


# ---------------------------------------------------------------- API
_RENDERERS = {"xlsx": to_xlsx, "docx": to_docx, "pdf": to_pdf}


def render(doc: Doc, fmt: str) -> Rendered:
    if fmt not in _RENDERERS:
        raise ExportError("Неизвестный формат выгрузки: допустимы xlsx, docx, pdf")
    return Rendered(_RENDERERS[fmt](doc), MEDIA_TYPES[fmt], f"{doc.filename}.{fmt}")


def content_disposition(filename: str) -> str:
    ascii_name = re.sub(r"[^A-Za-z0-9._-]+", "_", filename).strip("_") or "document"
    return f"attachment; filename=\"{ascii_name}\"; filename*=UTF-8''{quote(filename)}"
